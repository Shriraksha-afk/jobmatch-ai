from io import BytesIO
from pathlib import Path

def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract text from PDF, DOCX, or TXT files. Scanned PDFs need OCR separately."""
    ext = Path(filename).suffix.lower()
    if ext == ".txt":
        for encoding in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                return file_bytes.decode(encoding)
            except UnicodeDecodeError:
                pass
        return file_bytes.decode("utf-8", errors="replace")
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if ext == ".docx":
        from docx import Document
        doc = Document(BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs)
    raise ValueError("Unsupported file type. Please upload PDF, DOCX, or TXT.")
